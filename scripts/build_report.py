from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parent.parent
METRICS = PROJECT / "results" / "metrics.json"
OUTPUT = PROJECT / "deliverables" / "Atigini-Tani-Proje-Raporu-Gelistirilmis-v3.docx"
FIGURES = PROJECT / "figures"

NAVY = "173F3A"
TEAL = "2D6A63"
SAGE = "A7C4A0"
CORAL = "E76F51"
CREAM = "F5F0E6"
MIST = "EAF0EA"
WHITE = "FFFFFF"
GRAY = "66736F"
BLACK = "111111"
LINE = "D4DDD8"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 45, start: int = 90, bottom: int = 45, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_picture_alt(inline_shape, title_text: str, description: str) -> None:
    inline_shape._inline.docPr.set("title", title_text)
    inline_shape._inline.docPr.set("descr", description)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    run_properties.extend([color, underline, size])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([run_properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = BLACK, size: float = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def style_table(table, header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if header:
        mark_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            border(cell)
            if header and row_index == 0:
                shade(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                    run.font.bold = True
            elif row_index % 2 == 0:
                shade(cell, MIST)


def title(doc: Document, text: str, subtitle: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    if subtitle:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(9)
        run = paragraph.add_run(subtitle)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(15 if level == 1 else 11.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)


def body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        lead.font.name = "Arial"
        lead.font.size = Pt(9.5)
        text = text[len(bold_lead):]
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)


def bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Cm(0.55)
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
    if not paragraph.runs:
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
    else:
        paragraph.runs[0].text = text


def metric_cards(doc: Document, cards: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    width = Inches(6.95 / len(cards))
    for index, (value, label) in enumerate(cards):
        cell = table.cell(0, index)
        cell.width = width
        shade(cell, MIST)
        border(cell, SAGE, "8")
        set_cell_margins(cell, top=120, start=90, bottom=120, end=90)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(value)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        paragraph.add_run("\n")
        run = paragraph.add_run(label)
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Atığını Tanı  |  Makine ve Derin Öğrenme Final Projesi  |  ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def pct(value: float) -> str:
    return f"%{value * 100:.1f}".replace(".", ",")


def build() -> None:
    metrics = json.loads(METRICS.read_text(encoding="utf-8"))
    test = metrics["test"]
    model = metrics["model"]
    split = metrics["split"]
    report = metrics["classification_report"]
    matrix = metrics["confusion_matrix"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    # Sayfa 1
    title(doc, "Atığını Tanı", "Gerçek ortam atık görüntülerini dokuz malzeme sınıfında ayıran karar destek modeli")
    metric_cards(
        doc,
        [
            ("4.752", "gerçek ortam görüntüsü"),
            ("9", "malzeme sınıfı"),
            (pct(test["macro_f1"]), "makro F1"),
            (pct(test["accuracy"]), "test doğruluğu"),
        ],
    )
    heading(doc, "1. Problem ve gerçek dünya ihtiyacı")
    body(
        doc,
        "Bu proje, tek bir RGB atık görüntüsünü dokuz malzeme sınıfından birine atayan bir görüntü sınıflandırma modeli geliştirir. Girdi bir görüntüdür; hedef çıktı Cardboard, Food Organics, Glass, Metal, Miscellaneous Trash, Paper, Plastic, Textile Trash veya Vegetation etiketidir. Model, gerçek ayrıştırma hatlarında malzeme kararını destekleyebilir; insan denetiminin ve yerel ayrıştırma kurallarının yerine geçmez.",
    )
    heading(doc, "2. Veri seti")
    body(
        doc,
        "RealWaste, gerçek bir depolama sahasında çekilmiş 524 x 524 piksellik 4.752 renkli görüntü içerir. DOI: 10.24432/C5SS4G. UCI atıf kaydı CC BY 4.0, yazarın resmî GitHub README dosyası CC BY-NC-SA 4.0 belirtir. Bu çalışma daha kısıtlayıcı koşulu esas alır ve ham görüntüleri teslim paketine dahil etmez. Veri setinde eksik etiket yoktur; en büyük Plastic sınıfı ile en küçük Textile Trash sınıfı arasında yaklaşık 2,9 kat fark bulunur.",
    )
    table = doc.add_table(rows=4, cols=3)
    data = [
        ["Alan", "Değer", "Proje kararı"],
        ["Kaynak", "UCI + yazar deposu", "Kalıcı DOI ve lisans denetimi"],
        ["Bölme", f"{split['train']:,} / {split['validation']:,} / {split['test']:,}".replace(",", "."), "Katmanlı %70 / %15 / %15"],
        ["Dengesizlik", "318-921 örnek/sınıf", "Sınıf ağırlıklı çapraz entropi"],
    ]
    for row_index, row in enumerate(data):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value, bold=row_index == 0, color=WHITE if row_index == 0 else BLACK)
    style_table(table)
    heading(doc, "Temel sınır")
    body(doc, "Veri sahaya yakın olsa da tek bir toplama ortamını temsil eder. Farklı şehir, kamera, mevsim ve sıkıştırılmış/örtüşen atık koşullarında ek doğrulama gerekir.")

    # Sayfa 2
    page_break(doc)
    title(doc, "Veri keşfi ve ön işleme", "Sınıf dengesi, görsel çeşitlilik ve aydınlatma farklarının incelenmesi")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    distribution_picture = table.cell(0, 0).paragraphs[0].add_run().add_picture(
        str(FIGURES / "01-sinif-dagilimi.png"), width=Inches(3.35)
    )
    set_picture_alt(
        distribution_picture,
        "RealWaste sınıf dağılımı",
        "Dokuz atık sınıfındaki görüntü sayılarını gösteren sütun grafik",
    )
    sample_picture = table.cell(0, 1).paragraphs[0].add_run().add_picture(
        str(FIGURES / "02-sinif-ornekleri.png"), width=Inches(3.35)
    )
    set_picture_alt(
        sample_picture,
        "RealWaste sınıf örnekleri",
        "Dokuz atık sınıfının her birinden bir gerçek ortam görüntüsü",
    )
    for cell in table.rows[0].cells:
        border(cell, WHITE, "0")
    heading(doc, "3. Keşif bulguları")
    bullet(doc, "Plastic 921 örnekle en büyük; Textile Trash 318 örnekle en küçük sınıftır.")
    bullet(doc, "Gerçek saha arka planları temiz ürün fotoğraflarına göre daha karmaşıktır; örtüşme ve kirlenme sınıf ipuçlarını zayıflatır.")
    bullet(doc, "Sınıfların parlaklık dağılımları farklıdır; model yalnızca renge güvenmemeli, biçim ve doku örüntülerini de öğrenmelidir.")
    heading(doc, "4. Ön işleme ve veri sızıntısı kontrolü")
    steps = [
        ["Adım", "Uygulama", "Gerekçe"],
        ["Bölme", "Katmanlı %70/%15/%15, seed 42", "Sınıf oranlarını korur"],
        ["Boyut", "224 x 224 RGB", "Ön eğitim çözünürlüğüyle uyum"],
        ["Normalizasyon", "Her kanalda [-1, 1]", "MobileNetV2 ön işlemesiyle uyum"],
        ["Artırma", "Yatay çevirme, +/-12 derece dönüş, hafif parlaklık/kontrast", "Yalnızca eğitim verisi; saha çeşitliliği"],
        ["Denge", "Sınıf ağırlıklı kayıp", "Azınlık sınıflarını görünür kılar"],
    ]
    table = doc.add_table(rows=len(steps), cols=3)
    for row_index, row in enumerate(steps):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value, bold=row_index == 0, color=WHITE if row_index == 0 else BLACK, size=8.3)
    style_table(table)
    audit = metrics["reproducibility"]["exact_duplicate_audit"]
    body(doc, f"Bölme öncesi {audit['files_checked']:,} dosya SHA-256 ile denetlendi; tam içerik kopyası bulunmadı. Veri artırma yalnızca eğitim veri yükleyicisinde uygulandığı için doğrulama ve test setine yapay türev taşınmadı.".replace(f"{audit['files_checked']:,}", f"{audit['files_checked']:,}".replace(",", ".")))

    # Sayfa 3
    page_break(doc)
    title(doc, "Model tasarımı ve eğitim", "Ön eğitimli ve taşınabilir bir mimariyle yeniden üretilebilir eğitim akışı")
    heading(doc, "5. Mimari")
    architecture = [
        ["Bileşen", "Yapı", "İşlev"],
        ["Girdi", "224 x 224 x 3", "RGB görüntü"],
        ["Başlangıç", "ImageNet-1K ön eğitimi", "Genel görsel özellikleri aktarır"],
        ["Omurga", "MobileNetV2 ters artık blokları", "Çok ölçekli biçim ve doku"],
        ["Evrişim", "Derinlik ayrılabilir evrişim", "Düşük parametre ve işlem maliyeti"],
        ["Özet", "Global ortalama havuzlama", "Uzamsal özellikleri özetler"],
        ["Çıkış", "Yeni 9 sınıflı katman", "Sınıf ağırlıklı çapraz entropi"],
    ]
    table = doc.add_table(rows=len(architecture), cols=3)
    for row_index, row in enumerate(architecture):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value, bold=row_index == 0, color=WHITE if row_index == 0 else BLACK, size=8.2)
    style_table(table)
    heading(doc, "6. Eğitim ayarları")
    body(
        doc,
        f"Model Apple MPS üzerinde, mini-batch boyutu {model['batch_size']} ve {model['learning_rate']:.0e} başlangıç öğrenme oranıyla AdamW kullanılarak uçtan uca ince ayarlandı. Doğrulama kaybı izlenerek {model['epochs_completed']} epoch içindeki en iyi ağırlık saklandı; iyileşme durduğunda öğrenme oranını azaltma ve erken durdurma korumaları tanımlandı. Toplam parametre sayısı {model['parameters']:,}; deney seed 42 ile yeniden üretilebilir biçimde yürütüldü.".replace(f"{model['parameters']:,}", f"{model['parameters']:,}".replace(",", ".")),
    )
    training_picture = doc.add_picture(str(FIGURES / "04-egitim-egrileri.png"), width=Inches(6.75))
    set_picture_alt(
        training_picture,
        "Eğitim ve doğrulama eğrileri",
        "Dört epoch boyunca eğitim ve doğrulama kaybı ile doğruluk değerleri",
    )
    body(doc, "Model seçimi gerekçesi: MobileNetV2, küçük veri kümelerinde ImageNet görsel bilgisini aktarırken ters artık blokları ve derinlik ayrılabilir evrişimlerle taşınabilir bir yapı sunar. Dokuz sınıflı son katman yeniden öğrenilmiş, omurganın tamamı düşük öğrenme oranıyla ince ayarlanmıştır.")

    # Sayfa 4
    page_break(doc)
    title(doc, "Değerlendirme sonuçları", "Genel başarı, sınıf dengesi ve hata maliyeti birlikte değerlendirildi")
    metric_cards(
        doc,
        [
            (pct(test["accuracy"]), "doğruluk"),
            (pct(test["macro_precision"]), "makro kesinlik"),
            (pct(test["macro_recall"]), "makro duyarlılık"),
            (pct(test["macro_f1"]), "makro F1"),
            (pct(test["top2_accuracy"]), "ilk 2 doğruluk"),
        ],
    )
    body(
        doc,
        f"Her görüntüye çoğunluk sınıfını atayan temel yaklaşımın doğruluğu {pct(test['majority_baseline_accuracy'])}; modelin doğruluğu {pct(test['accuracy'])} ve makro F1 değeri {pct(test['macro_f1'])} oldu. Makro F1, küçük sınıfların büyük sınıflar tarafından gölgelenmesini engellediği için ana değerlendirme ölçütüdür.",
    )
    matrix_picture = doc.add_picture(str(FIGURES / "05-confusion-matrix.png"), width=Inches(5.55))
    set_picture_alt(
        matrix_picture,
        "Normalize karmaşıklık matrisi",
        "Dokuz sınıf için gerçek ve tahmin edilen etiketlerin normalize karmaşıklık matrisi",
    )
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(doc, "Sınıf bazlı F1")
    f1_rows = [["Sınıf", "Kesinlik", "Duyarlılık", "F1"]]
    for class_name in metrics["dataset"]["classes"]:
        values = report[class_name]
        f1_rows.append([class_name, pct(values["precision"]), pct(values["recall"]), pct(values["f1-score"])])
    table = doc.add_table(rows=len(f1_rows), cols=4)
    for row_index, row in enumerate(f1_rows):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value, bold=row_index == 0, color=WHITE if row_index == 0 else BLACK, size=7.7)
    style_table(table)

    # Sayfa 5
    page_break(doc)
    title(doc, "Hata analizi, sonuç ve geliştirme", "Yanlış tahmin örnekleri ve geliştirme öncelikleri")
    error_picture = doc.add_picture(str(FIGURES / "06-hata-galerisi.png"), width=Inches(6.45))
    set_picture_alt(
        error_picture,
        "Yüksek güvenli yanlış tahminler",
        "Modelin yüksek güvenle yanlış sınıflandırdığı on iki test görüntüsü",
    )
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(doc, "7. Hata örüntüleri")
    off_diagonal = []
    for true_index, row in enumerate(matrix):
        for pred_index, count in enumerate(row):
            if true_index != pred_index and count:
                off_diagonal.append((count, metrics["dataset"]["classes"][true_index], metrics["dataset"]["classes"][pred_index]))
    off_diagonal.sort(reverse=True)
    for count, true_name, pred_name in off_diagonal[:3]:
        bullet(doc, f"{true_name} → {pred_name}: {count} test görüntüsü. Benzer doku, kirlenme veya örtüşen nesneler olası etkendir.")
    heading(doc, "8. Güçlü yönler, sınırlar ve sonraki adımlar")
    bullet(doc, "Güçlü yön: Gerçek ortam verisi, katmanlı ayrım, sınıf ağırlığı ve ayrı test setiyle dürüst ölçüm.")
    bullet(doc, "Sınır: Tek veri kaynağı; coğrafya, kamera ve yerel ayrıştırma kuralı değişimlerine karşı dış doğrulama yok.")
    bullet(doc, "Geliştirme: Daha çeşitli saha verisi, etiket denetimi, ConvNeXt/EfficientNet karşılaştırması ve Grad-CAM ile karar bölgelerinin incelenmesi.")
    bullet(doc, "Ürünleşme: Güveni düşük tahminleri insana yönlendiren eşik, veri kayması takibi ve sınıf bazlı hata izleme.")
    heading(doc, "Sonuç")
    body(doc, f"Atığını Tanı, dokuz sınıflı gerçek ortam probleminde çoğunluk sınıfı temel yaklaşımının belirgin üzerinde ve makro F1 {pct(test['macro_f1'])} düzeyinde bir karar desteği üretmiştir. Hata örnekleri ve dış doğrulama eksikliği nedeniyle model doğrudan tesis otomasyonuna hazır değildir.")
    body(doc, "Kaynak: UCI RealWaste - https://archive.ics.uci.edu/dataset/908/realwaste - DOI 10.24432/C5SS4G. Lisans çelişkisi nedeniyle yazar deposundaki CC BY-NC-SA 4.0 koşulları esas alınmıştır.")
    links = doc.add_paragraph()
    links.paragraph_format.space_after = Pt(0)
    lead = links.add_run("Proje bağlantıları  ")
    lead.bold = True
    lead.font.name = "Arial"
    lead.font.size = Pt(8.5)
    add_hyperlink(links, "GitHub proje deposu", "https://github.com/serratanis/atigini-tani-atik-siniflandirma")
    separator = links.add_run("  |  ")
    separator.font.name = "Arial"
    separator.font.size = Pt(8.5)
    add_hyperlink(links, "Google Colab çalışma defteri", "https://colab.research.google.com/github/serratanis/atigini-tani-atik-siniflandirma/blob/main/notebooks/Atigini_Tani_RealWaste.ipynb")

    doc.core_properties.title = "Atığını Tanı Gerçek Ortam Atık Sınıflandırma Proje Raporu"
    doc.core_properties.subject = "Makine ve Derin Öğrenme Uygulamaları Final Projesi"
    doc.core_properties.author = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
